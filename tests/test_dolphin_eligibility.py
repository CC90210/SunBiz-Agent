from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from scrubber.scoring import load_config  # noqa: E402
from scrubber.telegram_bridge import format_packet, send_deal  # noqa: E402
from scrubber.uw_scoring import score_uw_deal  # noqa: E402


def _deal(**overrides):
    deal = {
        "true_revenue_monthly": 100_000,
        "leverage_pct": 10,
        "position_count": 2,
        "industry": "Professional Services",
        "iso_broker": "Links Capital Group",
        "data_merge_notes": "Clean",
        "previously_submitted": True,
        "state": "Florida",
        "uw_account_count": 1,
        "monthly_underwriting": [
            {"account_number": 1, "month": "June", "true_revenue": 100_000, "leverage_pct": 10},
        ],
        "counted_funders": [
            {"funder": "Generic Capital", "payoff_amount": 20_000},
            {"funder": "Another Funder", "payoff_amount": None},
        ],
    }
    deal.update(overrides)
    return deal


def test_dolphin_blocks_every_nationwide_iso_variant() -> None:
    cfg = load_config()
    for iso in ("Nationwide", "Nationwide Advance", "The Nationwide ISO Shop"):
        result = score_uw_deal(_deal(iso_broker=iso), cfg)
        assert result["tier"] == "bad"
        assert "blocked ISO/broker" in result["decline_reason"]


def test_dolphin_requires_two_positions_unless_previously_submitted() -> None:
    cfg = load_config()
    for positions in (None, 0, 1):
        result = score_uw_deal(_deal(position_count=positions, previously_submitted=False), cfg)
        assert result["tier"] == "bad"
        assert "active lender positions" in result["decline_reason"]
    assert score_uw_deal(_deal(position_count=2), cfg)["tier"] == "good"
    assert score_uw_deal(_deal(position_count=1, previously_submitted=True), cfg)["tier"] == "good"


def test_dolphin_blocks_restricted_states_and_more_than_five_positions() -> None:
    cfg = load_config()
    for state in ("Texas", "UT", "Virginia", "VA"):
        result = score_uw_deal(_deal(state=state), cfg)
        assert result["tier"] == "bad"
        assert "restricted state" in result["decline_reason"]
    result = score_uw_deal(_deal(position_count=6), cfg)
    assert result["tier"] == "bad"
    assert "active lender positions 6 > 5" in result["decline_reason"]


def test_dolphin_uses_sheet_monthly_leverage_and_requires_under_40() -> None:
    cfg = load_config()
    for leverage in (40, 41):
        result = score_uw_deal(_deal(sheet_monthly_leverage=leverage, leverage_pct=10), cfg)
        assert result["tier"] == "bad"
        assert "monthly leverage" in result["decline_reason"]
    assert score_uw_deal(_deal(sheet_monthly_leverage=39.99, leverage_pct=80), cfg)["tier"] == "good"


def test_dolphin_blocks_known_payoff_below_15000_but_allows_blank() -> None:
    cfg = load_config()
    low = _deal(counted_funders=[{"funder": "Generic", "payoff_amount": 14_999}])
    result = score_uw_deal(low, cfg)
    assert result["tier"] == "bad"
    assert "payoff amount" in result["decline_reason"]
    blank = _deal(counted_funders=[{"funder": "Generic", "payoff_amount": None}])
    assert score_uw_deal(blank, cfg)["tier"] == "good"


def test_preferred_funder_forces_review_except_nationwide() -> None:
    cfg = load_config()
    names = [
        "DLP", "CFG", "CFG MS", "FDM", "Forward Financing", "Square Advance",
        "Overton Funding", "Flow Capital", "Can Capital", "Capitas", "Legend",
        "MCA Servicing",
    ]
    for name in names:
        preferred = _deal(
            state="Texas", position_count=1, previously_submitted=False,
            counted_funders=[{"funder": name, "payoff_amount": 5_000}],
        )
        assert score_uw_deal(preferred, cfg)["tier"] in {"good", "review"}
    preferred = _deal(counted_funders=[{"funder": "DLP", "payoff_amount": 5_000}])
    preferred["iso_broker"] = "Nationwide Advance"
    result = score_uw_deal(preferred, cfg)
    assert result["tier"] == "bad"
    assert "blocked ISO/broker" in result["decline_reason"]


def test_telegram_boundary_blocks_stale_ineligible_candidate(monkeypatch) -> None:
    called = False

    def fake_api(*_args, **_kwargs):
        nonlocal called
        called = True
        return {"ok": True}

    monkeypatch.setattr("scrubber.telegram_bridge.api", fake_api)
    candidate = {"lead_data": _deal(position_count=1, previously_submitted=False), "tier": "good", "score": 99}
    result = send_deal({"EZRA_TELEGRAM_CHAT_ID": "123"}, candidate, "candidate-id")
    assert not result["ok"]
    assert "active lender positions 1 < 2" in result["error"]
    assert not called


def test_telegram_packet_shows_funder_date_and_payoff_numbers() -> None:
    deal = _deal()
    deal["uw_all_positions"] = [{
        "funder": "Generic Capital", "cadence": "weekly", "paid_off": False,
        "leverage_pct": 12.5, "date_funded": "2026-06-01", "payoff_amount": 20_000,
    }]
    packet = format_packet({"lead_data": deal, "tier": "good", "score": 90})
    assert "funded 2026-06-01" in packet
    assert "payoff $20,000" in packet


def test_dolphin_rejects_bad_month_hidden_by_good_average() -> None:
    cfg = load_config()
    result = score_uw_deal(_deal(
        true_revenue_monthly=346_667,
        sheet_monthly_leverage=25,
        uw_account_count=1,
        monthly_underwriting=[
            {"account_number": 1, "month": "April", "true_revenue": 500_000, "leverage_pct": 10},
            {"account_number": 1, "month": "May", "true_revenue": 500_000, "leverage_pct": 12},
            {"account_number": 1, "month": "June", "true_revenue": 40_000, "leverage_pct": 60},
        ],
    ), cfg)
    assert result["tier"] == "bad"
    assert "June true revenue $40,000" in result["decline_reason"]
    assert "June leverage 60%" in result["decline_reason"]


def test_dolphin_allows_workable_monthly_range_and_max_two_accounts() -> None:
    cfg = load_config()
    rows = [
        {"account_number": 1, "month": "April", "true_revenue": 400_000, "leverage_pct": 20},
        {"account_number": 1, "month": "May", "true_revenue": 600_000, "leverage_pct": 15},
        {"account_number": 1, "month": "June", "true_revenue": 300_000, "leverage_pct": 30},
        {"account_number": 2, "month": "June", "true_revenue": 100_000, "leverage_pct": 10},
    ]
    result = score_uw_deal(_deal(uw_account_count=2, monthly_underwriting=rows), cfg)
    assert result["tier"] == "good"
    assert any("monthly UW:" in reason for reason in result["reasons"])


def test_dolphin_rejects_more_than_two_uw_accounts_even_with_preferred_funder() -> None:
    cfg = load_config()
    result = score_uw_deal(_deal(
        uw_account_count=3,
        counted_funders=[{"funder": "DLP", "payoff_amount": 20_000}],
    ), cfg)
    assert result["tier"] == "bad"
    assert "business bank accounts 3 > 2" in result["decline_reason"]


def test_dolphin_blocks_stale_candidate_without_monthly_uw_evidence() -> None:
    cfg = load_config()
    result = score_uw_deal(_deal(uw_account_count=None, monthly_underwriting=[]), cfg)
    assert result["tier"] == "bad"
    assert "monthly revenue tables missing or unreadable" in result["decline_reason"]


def test_dolphin_requires_monthly_evidence_for_every_counted_account() -> None:
    cfg = load_config()
    result = score_uw_deal(_deal(uw_account_count=2), cfg)
    assert result["tier"] == "bad"
    assert "readable for 1 of 2 account(s)" in result["decline_reason"]


def test_dolphin_requires_both_revenue_and_leverage_for_every_month() -> None:
    cfg = load_config()
    result = score_uw_deal(_deal(monthly_underwriting=[
        {"account_number": 1, "month": "June", "true_revenue": 100_000, "leverage_pct": None},
    ]), cfg)
    assert result["tier"] == "bad"
    assert "June missing or unreadable leverage" in result["decline_reason"]


def test_parser_reads_every_month_and_counts_repeated_uw_tables() -> None:
    import openpyxl
    from scrubber.uw_sheet_parser import parse_uw_sheet

    wb = openpyxl.Workbook()
    ws = wb.active
    for start, values in ((10, (500_000, 500_000, 40_000)), (20, (100_000, 120_000, 110_000))):
        ws.cell(start, 2, "Month")
        ws.cell(start, 3, "True Revenue")
        ws.cell(start, 4, "Monthly Leverage")
        for offset, (month, revenue) in enumerate(zip(("April", "May", "June"), values), 1):
            ws.cell(start + offset, 2, month)
            ws.cell(start + offset, 3, revenue)
            ws.cell(start + offset, 4, (10 + offset) / 100)
        ws.cell(start + 4, 2, "Average")
        ws.cell(start + 4, 3, sum(values) / 3)
        ws.cell(start + 4, 4, 0.12)

    parsed = parse_uw_sheet(wb)
    assert parsed["uw_account_count"] == 2
    assert len(parsed["monthly_underwriting"]) == 6
    assert parsed["monthly_underwriting"][2] == {
        "month": "June", "true_revenue": 40_000.0,
        "leverage_pct": 13.0, "account_number": 1,
    }


def test_telegram_packet_shows_monthly_underwriting_context() -> None:
    deal = _deal(
        uw_account_count=1,
        monthly_underwriting=[
            {"account_number": 1, "month": "June", "true_revenue": 300_000, "leverage_pct": 30},
        ],
    )
    packet = format_packet({"lead_data": deal, "tier": "good", "score": 90})
    assert "UW accounts: 1" in packet
    assert "A1 June: $300,000 · 30% lev" in packet


def test_parser_pairs_side_by_side_account_columns() -> None:
    import openpyxl
    from scrubber.uw_sheet_parser import parse_uw_sheet

    wb = openpyxl.Workbook()
    ws = wb.active
    for col, revenue, leverage in ((1, 100_000, 0.10), (5, 200_000, 0.20)):
        ws.cell(10, col, "Month")
        ws.cell(10, col + 1, "True Revenue")
        ws.cell(10, col + 2, "Monthly Leverage")
        ws.cell(11, col, "June")
        ws.cell(11, col + 1, revenue)
        ws.cell(11, col + 2, leverage)
        ws.cell(12, col, "Average")
        ws.cell(12, col + 1, revenue)
        ws.cell(12, col + 2, leverage)

    rows = parse_uw_sheet(wb)["monthly_underwriting"]
    assert rows == [
        {"month": "June", "true_revenue": 100_000.0, "leverage_pct": 10.0, "account_number": 1},
        {"month": "June", "true_revenue": 200_000.0, "leverage_pct": 20.0, "account_number": 2},
    ]


def test_parser_preserves_fully_unreadable_labeled_month() -> None:
    import openpyxl
    from scrubber.uw_sheet_parser import parse_uw_sheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Month", "True Revenue", "Monthly Leverage"])
    ws.append(["June", "#VALUE!", "#DIV/0!"])
    ws.append(["Average", 100_000, 0.20])

    parsed = parse_uw_sheet(wb)
    assert parsed["monthly_underwriting"] == [{
        "month": "June", "true_revenue": None,
        "leverage_pct": None, "account_number": 1,
    }]


def test_parser_counts_table_with_missing_leverage_header() -> None:
    import openpyxl
    from scrubber.uw_sheet_parser import parse_uw_sheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Month", "True Revenue", "Monthly Leverage", "", "Month", "True Revenue"])
    ws.append(["June", 100_000, 0.10, "", "June", 200_000])
    ws.append(["Average", 100_000, 0.10, "", "Average", 200_000])

    parsed = parse_uw_sheet(wb)
    assert parsed["uw_account_count"] == 2
    assert parsed["uw_revenue_tables"][1]["parse_error"] == "Monthly Leverage header missing or unreadable"
    result = score_uw_deal({**_deal(), **parsed}, load_config())
    assert result["tier"] == "bad"
    assert "readable for 1 of 2 account(s)" in result["decline_reason"]
